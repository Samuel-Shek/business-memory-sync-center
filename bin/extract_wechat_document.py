#!/usr/bin/env python3
"""Safely extract passive text from a single local business document."""

from __future__ import annotations

from archive_paths import path as _archive_path
from archive_guard import require_opt_in as _archive_require_opt_in
_archive_require_opt_in()


import argparse
import json
import os
import sys
import zipfile
from pathlib import Path


os.umask(0o077)
MAX_TEXT = 2_000_000
MAX_CELLS = 100_000
MAX_PAGES = 5_000
MAX_ARCHIVE_ENTRIES = 20_000
MAX_UNCOMPRESSED_BYTES = 1024 * 1024 * 1024
FORBIDDEN_PACKAGE_SUFFIXES = {
    ".exe", ".dll", ".dylib", ".app", ".js", ".vbs", ".bat", ".cmd", ".ps1",
    ".sh", ".py", ".jar", ".com", ".msi", ".pkg", ".dmg",
}


def arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("path", type=Path)
    return parser.parse_args()


def trim(parts: list[str]) -> tuple[str, bool]:
    text = "\n".join(part for part in parts if part).strip()
    if len(text) <= MAX_TEXT:
        return text, False
    return text[:MAX_TEXT] + "\n[文档文字过长，已截断]", True


def validate_ooxml(path: Path) -> None:
    if not zipfile.is_zipfile(path):
        raise ValueError("ooxml_invalid")
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_ARCHIVE_ENTRIES:
            raise ValueError("archive_entry_limit_exceeded")
        total = 0
        for item in entries:
            normalized = item.filename.replace("\\", "/").rstrip("/")
            parts = normalized.split("/")
            if not normalized or normalized.startswith("/") or any(part in {"", ".", ".."} for part in parts):
                raise ValueError("archive_path_invalid")
            if item.flag_bits & 0x1:
                raise ValueError("archive_encrypted")
            if ((item.external_attr >> 16) & 0o170000) == 0o120000:
                raise ValueError("archive_path_invalid")
            total += item.file_size
            if total > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("archive_size_limit_exceeded")
            lower = normalized.casefold()
            if lower.endswith("vbaproject.bin") or Path(lower).suffix in FORBIDDEN_PACKAGE_SUFFIXES:
                raise ValueError("active_content_blocked")


def extract_pdf(path: Path) -> dict:
    from pypdf import PdfReader

    reader = PdfReader(str(path), strict=False)
    if reader.is_encrypted:
        raise ValueError("archive_encrypted")
    catalog = reader.trailer.get("/Root")
    catalog = catalog.get_object() if hasattr(catalog, "get_object") else catalog
    if isinstance(catalog, dict):
        if catalog.get("/OpenAction") is not None or catalog.get("/AA") is not None:
            raise ValueError("active_content_blocked")
        names = catalog.get("/Names")
        names = names.get_object() if hasattr(names, "get_object") else names
        if isinstance(names, dict) and any(key in names for key in ("/JavaScript", "/EmbeddedFiles")):
            raise ValueError("active_content_blocked")
    if len(reader.pages) > MAX_PAGES:
        raise ValueError("page_limit_exceeded")
    text, truncated = trim([(page.extract_text() or "") for page in reader.pages])
    return {"kind": "pdf", "pages": len(reader.pages), "text": text, "truncated": truncated}


def extract_docx(path: Path) -> dict:
    from docx import Document

    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    cells = 0
    for table in document.tables:
        for row in table.rows:
            values = [cell.text for cell in row.cells]
            cells += len(values)
            if cells > MAX_CELLS:
                raise ValueError("cell_limit_exceeded")
            parts.append("\t".join(values))
    text, truncated = trim(parts)
    return {"kind": "word", "paragraphs": len(document.paragraphs), "cells": cells, "text": text, "truncated": truncated}


def extract_xlsx(path: Path) -> dict:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True, keep_links=False)
    parts: list[str] = []
    cells = 0
    try:
        for sheet in workbook.worksheets:
            parts.append(f"[工作表] {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                cells += len(values)
                if cells > MAX_CELLS:
                    raise ValueError("cell_limit_exceeded")
                if any(values):
                    parts.append("\t".join(values))
    finally:
        workbook.close()
    text, truncated = trim(parts)
    return {"kind": "excel", "sheets": len(workbook.sheetnames), "cells": cells, "text": text, "truncated": truncated}


def extract_pptx(path: Path) -> dict:
    from pptx import Presentation

    presentation = Presentation(str(path))
    parts: list[str] = []
    for index, slide in enumerate(presentation.slides, 1):
        parts.append(f"[幻灯片 {index}]")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    parts.append("\t".join(cell.text for cell in row.cells))
    text, truncated = trim(parts)
    return {"kind": "powerpoint", "slides": len(presentation.slides), "text": text, "truncated": truncated}


def main() -> int:
    args = arguments()
    path = args.path.resolve()
    if not path.is_file() or path.is_symlink() or path.stat().st_size > 512 * 1024 * 1024:
        print(json.dumps({"status": "failed", "error_code": "document_invalid"}))
        return 2
    extension = path.suffix.casefold()
    try:
        if extension == ".pdf":
            result = extract_pdf(path)
        elif extension == ".docx":
            validate_ooxml(path)
            result = extract_docx(path)
        elif extension == ".xlsx":
            validate_ooxml(path)
            result = extract_xlsx(path)
        elif extension == ".pptx":
            validate_ooxml(path)
            result = extract_pptx(path)
        else:
            print(json.dumps({"status": "failed", "error_code": "document_type_unsupported"}))
            return 3
        print(json.dumps({"status": "success", **result}, ensure_ascii=False, sort_keys=True))
        return 0
    except Exception as exc:
        allowed = {
            "page_limit_exceeded", "cell_limit_exceeded", "ooxml_invalid",
            "archive_entry_limit_exceeded", "archive_path_invalid", "archive_encrypted",
            "archive_size_limit_exceeded", "active_content_blocked",
        }
        code = str(exc) if str(exc) in allowed else "document_parse_failed"
        print(json.dumps({"status": "failed", "error_code": code}, ensure_ascii=False))
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
